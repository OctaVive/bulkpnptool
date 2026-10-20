"""End-to-end checks against the running PHP app, using Python as a reference."""
import io
from pathlib import Path
import sys
import unittest
from urllib.error import HTTPError
from urllib.parse import urlencode
from urllib.request import urlopen
import zipfile

sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
from app import app
from openpyxl import load_workbook

BASE = 'http://127.0.0.1:5001'
def fetch(path='/', data=None):
    try:
        response=urlopen(BASE+path,data=urlencode(data).encode() if data is not None else None)
    except HTTPError as error:
        response=error
    return response.status,response.headers,response.read()

class HttpTests(unittest.TestCase):
    def test_tabs(self):
        for tab in ('provisioning','splitter','configuration'):
            status,_,body=fetch('/?tab='+tab)
            self.assertEqual(status,200)
            self.assertIn(f'class="tab-panel active" data-tab-panel="{tab}"'.encode(),body)
            self.assertNotIn(b'Warning:',body)
            self.assertIn(b'Coming soon',body)

    def test_csv_downloads(self):
        for operation in ('ADD','DELETE','MOVE'):
            for numbers in ('0455688200\n45568820X\n+31711234567','0501234567\n0521123456'):
                data=dict(operation=operation,numbers=numbers,pbx_id_from='1000001',pbx_id_to='1000002',confirm_050='1',pe_0501234567='01',pe_0521123456='12')
                expected=app.test_client().post('/',data=data)
                status,headers,body=fetch(data=data)
                self.assertEqual(status,expected.status_code)
                self.assertEqual(body,expected.data)
                self.assertEqual(headers['Content-Disposition'],expected.headers['Content-Disposition'])

    def test_validation(self):
        for operation in ('ADD','DELETE','MOVE'):
            for field in (['pbx_id_from','pbx_id_to'] if operation=='MOVE' else ['pbx_id_to' if operation=='ADD' else 'pbx_id_from']):
                for invalid in ('','123456','123abc7'):
                    for fmt in ('csv','xlsx'):
                        data=dict(operation=operation,numbers='0455688200',pbx_id_from='1000001',pbx_id_to='1000002',output_format=fmt)
                        data[field]=invalid
                        status,headers,body=fetch(data=data)
                        self.assertEqual(status,400)
                        self.assertIn(b'must contain at least 7 digits',body)
                        self.assertNotIn('Content-Disposition',headers)
        for fmt in ('csv','xlsx'):
            status,_,body=fetch(data=dict(operation='MOVE',pbx_id_from='1000001',pbx_id_to=' 1000001 ',output_format=fmt))
            self.assertEqual(status,400)
            self.assertIn(b'You cannot move to the same PBX-ID',body)

    def test_location_selection(self):
        for numbers in ('0501234567','0501234567\n0521123456'):
            status,_,body=fetch(data=dict(operation='ADD',numbers=numbers,pbx_id_to='1000001'))
            self.assertEqual(status,200)
            self.assertIn(b'Select PE location for special prefixes',body)
            self.assertIn(b'name="pe_0501234567"',body)
            self.assertNotIn(b'Warning:',body)
            self.assertEqual(b'id="bulk_location"' in body,'\n' not in numbers)

    def test_xlsx(self):
        data=dict(operation='ADD',numbers='455688200\n45568820X\n+31501234567\n0031711234567',pbx_id_to='1000001',output_format='xlsx')
        expected=app.test_client().post('/',data=data)
        status,headers,body=fetch(data=data)
        self.assertEqual(status,200)
        actual_book=load_workbook(io.BytesIO(body))
        expected_book=load_workbook(io.BytesIO(expected.data))
        self.assertEqual(list(actual_book.active.values),list(expected_book.active.values))
        self.assertEqual(actual_book.active.title,'Blocks')
        self.assertEqual(actual_book.active['A1'].fill.fgColor.rgb,'FFE60000')
        self.assertTrue(actual_book.active['A1'].font.bold)

    def test_splitter(self):
        for sources,removes in [('04556882xx','0455688250'),('05012345xx','0501234500'),('invalid','bad')]:
            data=dict(sources=sources,removes=removes,splitter_pbx_id='1000001')
            for path in ('/splitter/calculate','/splitter/export-text','/splitter/download'):
                with self.subTest(sources=sources,path=path):
                    expected=app.test_client().post(path,data=data)
                    status,headers,body=fetch(path,data)
                    self.assertEqual(status,expected.status_code)
                    self.assertNotIn(b'Warning:',body)
                    if path.endswith('export-text'): self.assertEqual(body,expected.data)
                    elif headers.get_content_type()=='application/zip':
                        with zipfile.ZipFile(io.BytesIO(body)) as actual, zipfile.ZipFile(io.BytesIO(expected.data)) as reference:
                            self.assertEqual(actual.namelist(),reference.namelist())
                            for name in actual.namelist(): self.assertEqual(actual.read(name),reference.read(name))
                    else: self.assertIn(b'class="tab-panel active" data-tab-panel="splitter"',body)
        data=dict(sources='05012345xx',removes='0501234500',splitter_pbx_id='1000001',spe_source_0501234500_100='01')
        status,headers,body=fetch('/splitter/download',data)
        self.assertEqual(status,200)
        self.assertEqual(headers.get_content_type(),'application/zip')
        self.assertIn('add1000001.csv',zipfile.ZipFile(io.BytesIO(body)).namelist())

    def test_assets_and_private_files(self):
        status,headers,body=fetch('/assets/Icon-88248495-6083-4dcf-8cb0-1d42d1fb34e3.png')
        self.assertEqual(status,200)
        self.assertEqual(body,(Path(__file__).resolve().parents[2]/'assets/Icon-88248495-6083-4dcf-8cb0-1d42d1fb34e3.png').read_bytes())
        for path in ('/lib.php','/data/pe_codes.txt','/tests/fixtures.json','/missing'):
            self.assertEqual(fetch(path)[0],404)

if __name__=='__main__': unittest.main()
